import re
from pathlib import Path

import docx


FILES = [
    Path(r"C:\Users\91916\Downloads\Ashwin's Blackbook.docx"),
    Path(r"C:\Users\91916\Downloads\Final Report Format.docx"),
]


def main():
    for file_path in FILES:
        print(f"\nFILE {file_path}")
        document = docx.Document(file_path)
        print("paragraphs", len(document.paragraphs), "tables", len(document.tables))
        for i, paragraph in enumerate(document.paragraphs):
            text = " ".join(paragraph.text.split())
            if not text:
                continue
            style = paragraph.style.name if paragraph.style else ""
            is_candidate = (
                style.startswith("Heading")
                or re.match(
                    r"^(CHAPTER|Chapter|\d+(\.\d+)+|ABSTRACT|REFERENCES|APPENDIX|CONCLUSIONS|SYSTEM|LITERATURE|METHODOLOGY|RISK|INTRODUCTION)",
                    text,
                    re.I,
                )
                or "XXXXX" in text
                or text.lower() == "bye"
            )
            if is_candidate:
                print(f"{i:04d}\t{style}\t{text[:240]}")

        for ti, table in enumerate(document.tables):
            print(f"TABLE {ti} rows={len(table.rows)} cols={len(table.columns)}")
            for ri, row in enumerate(table.rows[:4]):
                cells = " | ".join(
                    cell.text.strip().replace("\n", " / ")[:140] for cell in row.cells
                )
                print(f"  {ri}: {cells}")


if __name__ == "__main__":
    main()
