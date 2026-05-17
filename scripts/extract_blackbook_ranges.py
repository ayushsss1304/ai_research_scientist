from pathlib import Path
import sys

import docx


DOC = Path(r"C:\Users\91916\Downloads\Ashwin's Blackbook.docx")
RANGES = [(465, 514), (908, 1014), (1014, 1092), (1057, 1148)]


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    document = docx.Document(DOC)
    for start, end in RANGES:
        print(f"\n--- {start} {end} ---")
        for i in range(start, min(end, len(document.paragraphs))):
            text = " ".join(document.paragraphs[i].text.split())
            if text:
                print(f"{i:04d}: {text[:700]}")


if __name__ == "__main__":
    main()
