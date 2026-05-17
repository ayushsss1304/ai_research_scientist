from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.text.paragraph import Paragraph


FINAL_REPORT = Path(r"C:\Users\91916\Downloads\Final Report Format.docx")
OUTPUT = Path(
    r"C:\Users\91916\OneDrive\Desktop\Btech Project PY\research_scientist\docs\Final_Report_Completed_From_Blackbook_Reference.docx"
)


def normalize(text: str) -> str:
    return " ".join((text or "").split())


def paragraph_after(paragraph, text=""):
    new_p = deepcopy(paragraph._p)
    paragraph._p.addnext(new_p)
    inserted = Paragraph(new_p, paragraph._parent)
    inserted.clear()
    if text:
        inserted.add_run(text)
    return inserted


def find_paragraph(doc, exact):
    target = normalize(exact)
    for p in doc.paragraphs:
        if normalize(p.text) == target:
            return p
    raise ValueError(f"Could not find paragraph: {exact}")


def replace_paragraph(doc, exact, text):
    p = find_paragraph(doc, exact)
    p.clear()
    p.add_run(text)
    return p


def add_after(anchor, lines):
    current = anchor
    for line in lines:
        current = paragraph_after(current, line)
    return current


def replace_if_found(doc, exact, text):
    try:
        return replace_paragraph(doc, exact, text)
    except ValueError:
        return None


def main():
    doc = Document(FINAL_REPORT)

    # Targeted corrections so the existing report matches the current project.
    replace_if_found(
        doc,
        "To design and build an AI-based multi-agent research assistant that automates literature discovery and analysis providing accurate, personalized, and evidence-verified insights to accelerate scientific innovation.",
        "To design and build VeriSci, an AI-based research assistant that automates literature discovery, research-gap identification, legal open-access full-text enrichment, knowledge organization, and evidence-grounded question answering using Scopus/Elsevier, IEEE, open academic sources, RAG, and knowledge graph techniques.",
    )
    replace_if_found(
        doc,
        "Documentation and Research Paper Preparation: The team jointly contributed to report writing, conference paper preparation, plagiarism verification, and project presentations.",
        "Documentation and Research Paper Preparation: The team jointly contributed to report writing, conference paper preparation, result analysis, diagram preparation, and project presentations.",
    )
    replace_if_found(
        doc,
        "The project depends on multiple academic repositories such as ArXiv, IEEE, and Semantic Scholar.",
        "The project depends on multiple academic repositories and APIs such as Scopus/Elsevier, IEEE, Semantic Scholar, OpenAlex, ArXiv, CrossRef, PubMed, CORE, and Google Scholar.",
    )
    replace_if_found(
        doc,
        "The extracted text is divided into smaller, semantically coherent chunks (e.g., by paragraph or section). Each chunk is converted into a high-dimensional vector embedding using a pre-trained language model and indexed into a vector database (e.g., FAISS or ChromaDB) for efficient similarity retrieval.",
        "The extracted text is divided into smaller, semantically coherent chunks. Each chunk is converted into a high-dimensional vector embedding using a sentence-transformer model. In the current MVP, the embeddings and document chunks are stored in a lightweight local document store rather than a heavy external vector database, making the system easier to run on a laptop.",
    )
    replace_if_found(
        doc,
        "The VeriSci project uses two kinds of databases to handle parts of the system. It uses a Knowledge Graph Database, which is also known as Neo4j to deal with relationships and metadata. The VeriSci project also uses a Custom Local Vector/Document Store for something called Retrieval-Augmented Generation or RAG for short.",
        "The VeriSci project uses three storage mechanisms for different responsibilities. Neo4j is used as the knowledge graph database for relationships between papers, authors, and research fields. A custom local document/vector store is used for Retrieval-Augmented Generation (RAG). A lightweight SQLite database stores workspaces and saved papers so that users can continue literature review work across sessions.",
    )

    # 4.6 SOCME: Ashwin's report uses a SOCME chapter with core component interaction and applications.
    replace_if_found(
        doc,
        "bye",
        "The System-Oriented Concept Map Extension (SOCME) for VeriSci represents the project as a set of interacting subsystems rather than a simple linear workflow. Following the reference format, the system is divided into core components and their interactions are described through labelled relationships. The main subsystems are User Interaction, Literature Retrieval, Paper Access, AI Analysis, Knowledge Organization, and Output/Storage.",
    )
    socme_anchor = find_paragraph(
        doc,
        "The System-Oriented Concept Map Extension (SOCME) for VeriSci represents the project as a set of interacting subsystems rather than a simple linear workflow. Following the reference format, the system is divided into core components and their interactions are described through labelled relationships. The main subsystems are User Interaction, Literature Retrieval, Paper Access, AI Analysis, Knowledge Organization, and Output/Storage.",
    )
    add_after(
        socme_anchor,
        [
            "Core Components Interaction",
            "Researcher and User Interface: The researcher enters a topic, selects sources, manages workspaces, saves relevant papers, and requests research-gap analysis. The dashboard acts as the central interaction layer.",
            "Literature Retrieval Subsystem: The query is sent to verified sources such as Scopus/Elsevier and IEEE, and to discovery/open-access sources such as Semantic Scholar, OpenAlex, ArXiv, CrossRef, PubMed, CORE, and Google Scholar. The retrieved metadata includes titles, authors, abstracts, DOI/EID values, publication year, citations, journal/source details, and links.",
            "Metadata Normalization and Verification: Results from different APIs are converted into a common paper structure. Duplicate papers are removed using DOI matching, ArXiv identifiers, and fuzzy title comparison. DOI-bearing results are enriched with Scopus verification wherever available.",
            "Paper Access Subsystem: If the user selects Deep Research, the system checks for legal open-access PDFs through existing PDF links, OpenAlex, Unpaywall, or arXiv. Available PDFs are downloaded only into a temporary cache, useful sections are extracted, and the temporary files are deleted after processing.",
            "AI Analysis Subsystem: Abstracts, snippets, and extracted full-text sections are supplied to the RAG pipeline and research-gap generator. The system uses embeddings, similarity retrieval, and an LLM/Ollama chatbot to produce grounded answers and gap suggestions.",
            "Knowledge Organization Subsystem: Paper, Author, and Field nodes are created in the Neo4j knowledge graph. Relationships such as AUTHORED and IN_FIELD help users understand connections between publications, contributors, and research domains.",
            "Output and Storage Subsystem: SQLite stores workspaces and saved papers. The interface presents ranked search results, saved paper libraries, knowledge graph insights, research gap reports, and export-ready outputs.",
            "Fig. 4.5 SOCME Diagram for VeriSci Research Gap Finder",
            "Applications",
            "Academic Literature Review: VeriSci helps students and researchers search multiple academic sources, remove duplicate records, and identify relevant papers faster.",
            "Research Gap Identification: The system supports normal metadata-based gap analysis and Deep Research gap analysis using legal open-access full text where available.",
            "Project Report Preparation: Saved papers, source metadata, gap summaries, and export features support report writing and documentation.",
            "Knowledge Discovery: The knowledge graph allows users to explore author, paper, and field relationships for trend analysis.",
        ],
    )

    # Chapter 5 unfinished sections.
    add_after(
        find_paragraph(doc, "5.3 Important Algorithms"),
        [
            "5.3.1 Multi-Source Retrieval Algorithm",
            "The multi-source retrieval algorithm accepts a research query and selected filters from the user interface. It sends the query to selected sources, prioritizing verified academic databases such as Scopus and IEEE while also supporting open academic sources for broader discovery. Requests are executed in parallel to reduce latency.",
            "The returned records are normalized into a common paper dictionary. Duplicate records are removed using DOI matching, ArXiv ID comparison, and fuzzy title similarity. The final papers are ranked using relevance score, title match, keyword overlap, citation count, publication recency, source quality, and metadata completeness.",
            "5.3.2 Deep Research Gap Algorithm",
            "The normal gap finder uses titles, abstracts, keywords, citations, and source metadata to identify possible research gaps. Deep Research improves this by attempting to retrieve legal open-access PDFs. The system extracts compact evidence from sections such as abstract, methodology, results, limitations, conclusion, and future work. If full text is unavailable, the system falls back to metadata-only analysis.",
            "5.3.3 RAG Question Answering Algorithm",
            "Uploaded PDFs are processed by extracting text, cleaning it, dividing it into chunks, and converting each chunk into sentence-transformer embeddings. During question answering, the query embedding is compared with stored chunk embeddings using cosine similarity. Maximal Marginal Relevance is used to choose diverse and relevant chunks before sending context to the LLM.",
            "5.3.4 Knowledge Graph Construction Algorithm",
            "The knowledge graph construction algorithm converts paper metadata into Paper, Author, and Field nodes. Paper nodes store title, abstract, year, DOI, source, citations, journal, URL, and PDF availability. Authors are connected to papers using AUTHORED relationships, and papers are connected to domains through IN_FIELD relationships.",
        ],
    )

    add_after(
        find_paragraph(doc, "5.4 Important Data Structure"),
        [
            "The main data structures used in the VeriSci platform are described below.",
            "Paper Dictionary: A normalized dictionary containing title, authors, abstract, year, citations, source, journal, DOI, EID, URL, PDF URL, keywords, subject areas, Scopus verification status, and source metrics.",
            "Search Cache JSON: Search results are stored temporarily as JSON cache files using a cache key derived from query parameters and selected sources. This reduces response time for repeated searches.",
            "SQLite Workspace Tables: The workspaces table stores local research projects, while the saved_papers table stores selected papers, notes, tags, source details, DOI, year, citations, and the complete paper JSON.",
            "RAG Document Store: The document store maintains metadata, extracted text, chunks, and embeddings in a lightweight local serialized format. This avoids dependence on external vector databases for the MVP.",
            "Neo4j Graph Structure: Paper, Author, and Field nodes are linked through AUTHORED and IN_FIELD relationships. Indexing improves graph search by paper title, author name, and publication year.",
            "Deep Research Evidence Object: This structure stores the paper title, source, DOI/EID, evidence level, full-text status, extracted sections, evidence snippets, and the number of characters used for analysis.",
        ],
    )

    add_after(
        find_paragraph(doc, "5.5 Graphics User Interface Screenshot"),
        [
            "The user interface of VeriSci is organized as a dashboard-based web application. It allows users to search papers, manage workspaces, save selected papers, view paper details, generate research gaps, explore the knowledge graph, upload PDFs, and chat with uploaded documents.",
            "The following screenshots should be added in this section:",
            "Figure 5.1 Dashboard Workspace Screen",
            "Figure 5.2 Search Interface with Scopus/IEEE and Open-Access Source Selection",
            "Figure 5.3 Search Result Cards with Scopus Verification and PDF Badges",
            "Figure 5.4 Deep Research Gap Finder Output",
            "Figure 5.5 Knowledge Graph Explorer",
            "Figure 5.6 PDF Chat/RAG Interface",
        ],
    )

    # Chapter 7: mirror Ashwin's testing chapter format.
    add_after(
        find_paragraph(doc, "SYSTEM TESTING"),
        [
            "7.1 Type of Testing",
            "Testing was performed to verify the functional correctness, reliability, and integration of the VeriSci platform. The testing approach follows the reference blackbook structure and includes positive testing, negative testing, unit testing, and integration testing.",
            "7.1.1 Positive Testing",
            "Positive testing checks whether the system works correctly when valid inputs are provided. Test cases included searching valid research topics, selecting Scopus and IEEE sources, saving papers to a workspace, uploading readable PDFs, generating research gaps, exporting results, and viewing knowledge graph statistics. In all cases, the expected output was successful completion of the requested workflow with correct data display.",
            "7.1.2 Negative Testing",
            "Negative testing checks how the system behaves when invalid or incomplete inputs are given. Test cases included submitting an empty search query, uploading a non-PDF file, requesting RAG chat before indexing documents, selecting unavailable APIs, and running optional graph features when Neo4j is not active. The system handled these cases by showing validation messages or graceful error responses.",
            "7.1.3 Unit Testing",
            "Unit testing was performed on individual modules such as EnhancedPaperScraper, ScopusScraper, PDFProcessor, ResearchStore, RAG chatbot functions, export logic, and knowledge graph manager functions. Each unit was checked for expected input-output behavior, exception handling, and data consistency.",
            "7.1.4 Integration Testing",
            "Integration testing verified complete workflows across modules. The major workflows tested were search to save paper, search to knowledge graph insertion, PDF upload to RAG chat, search result to research gap generation, and Deep Research full-text extraction to evidence-based gap output.",
            "7.2 Testing Summary",
            "The testing process showed that the system can support the main research workflow from query search to paper saving, research gap analysis, knowledge graph organization, and PDF-based question answering. Optional services such as Ollama and Neo4j are handled separately, so the web application remains usable even if those services are temporarily unavailable.",
        ],
    )

    # Chapter 8: results.
    add_after(
        find_paragraph(doc, "EXPERIMENTAL RESULTS"),
        [
            "8.1 Results and Outputs",
            "The implemented VeriSci platform successfully provides a unified environment for academic paper discovery, organization, research-gap identification, and document-grounded question answering. The search module retrieves papers from verified sources such as Scopus/Elsevier and IEEE, while also supporting open academic sources for wider coverage.",
            "The result cards display title, authors, year, source, citations, DOI, journal, abstract, PDF availability, and Scopus verification status. This helps users quickly judge whether a paper is relevant and reliable. The workspace dashboard allows selected papers to be saved locally, making the literature review process more organized.",
            "The Deep Research feature improves the quality of research-gap generation by using legal open-access PDFs whenever available. It extracts compact evidence from useful sections of the paper and deletes the temporary PDF after processing. This reduces laptop storage load while allowing stronger evidence-based analysis.",
            "The RAG-based PDF chat module successfully answers user questions using uploaded research papers. The knowledge graph module organizes paper-author-field relationships and provides statistics and insights that support trend analysis.",
            "8.2 Performance Observations",
            "Parallel source searching reduces the delay caused by multiple API calls. Local caching improves repeated search performance because the system can reuse previous results instead of calling every external source again. OCR is disabled by default to keep normal PDF processing faster, but it can be enabled when scanned PDFs must be processed.",
            "8.3 Suggested Result Figures",
            "Figure 8.1 Search Results for a Research Topic",
            "Figure 8.2 Workspace Saved Paper Library",
            "Figure 8.3 Research Gap Finder Output",
            "Figure 8.4 Deep Research Evidence Output",
            "Figure 8.5 RAG Chat Response for Uploaded PDF",
            "Figure 8.6 Knowledge Graph Visualization",
            "8.4 Result Analysis",
            "The results indicate that VeriSci works as a practical MVP for AI-assisted literature review. Its main strength is combining verified metadata, open-access discovery, temporary full-text enrichment, local workspace storage, RAG-based document chat, and graph-based knowledge organization into one application. The main limitation is that full-text analysis depends on legal open-access availability; paywalled papers may provide only metadata and abstracts.",
        ],
    )

    # Chapter 9: conclusion and future scope.
    replace_if_found(doc, "Summary of what the project has been achieved.", "")
    replace_if_found(
        doc,
        "Must include your quantitative results and logical analysis of the result presented in the project report.",
        "",
    )
    add_after(
        find_paragraph(doc, "CONCLUSIONS AND FUTURE SCOPE"),
        [
            "9.1 Conclusion",
            "The VeriSci project successfully demonstrates an AI-based research assistance platform that supports the literature review lifecycle from paper discovery to research-gap identification. The system integrates Scopus/Elsevier, IEEE, and open academic sources to retrieve paper metadata, normalize results, remove duplicates, and rank papers according to relevance and quality.",
            "The platform improves research productivity through workspace-based storage, legal open-access full-text enrichment, RAG-based PDF chat, and knowledge graph visualization. Deep Research addresses the earlier limitation of title-only search by temporarily downloading legal open-access PDFs, extracting useful evidence, and deleting files after processing.",
            "Overall, VeriSci provides a practical MVP for students and researchers who need a unified tool for searching, organizing, analyzing, and understanding research papers. It shows how retrieval-based AI, local storage, knowledge graphs, and conversational interfaces can make academic research faster, more structured, and more evidence-driven.",
            "9.2 Future Scope",
            "1. Institutional Access Integration: The system can be extended to support university/library subscriptions for retrieving full text from subscribed sources in a legal manner.",
            "2. Advanced Knowledge Graph: Future versions can add Dataset, Method, Metric, Institution, Model, and Research Gap nodes to detect deeper research patterns.",
            "3. Automatic Literature Review Generation: The system can generate structured literature review drafts with comparison tables and citations.",
            "4. Collaboration Features: Multiple users can work in shared research workspaces with comments, roles, and version history.",
            "5. Citation Manager Integration: Zotero, Mendeley, or BibTeX synchronization can be added for smoother academic writing.",
            "6. Cloud Deployment: Deploying the application on a cloud platform can make it accessible from any device and support larger experiments.",
            "9.3 Applications",
            "Academic Research: Students and researchers can use VeriSci to discover papers, compare sources, and identify possible research gaps.",
            "Final-Year Project Work: The platform helps teams collect references, save important papers, prepare reports, and generate evidence-backed analysis.",
            "Research Supervision: Guides and faculty members can use saved workspaces and knowledge graph insights to review student literature coverage.",
            "Technical Writing: Exported metadata, summaries, and gap suggestions can support survey papers, project reports, and presentations.",
        ],
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
