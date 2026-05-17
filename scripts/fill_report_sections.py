from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.text.paragraph import Paragraph


SOURCE = Path(r"C:\Users\91916\Downloads\Final Report Format.docx")
OUT = Path(r"C:\Users\91916\OneDrive\Desktop\Btech Project PY\research_scientist\docs\Final_Report_Format_filled_VerISci.docx")


def insert_after(paragraph, text="", style=None):
    new_p = deepcopy(paragraph._p)
    paragraph._p.addnext(new_p)
    inserted = Paragraph(new_p, paragraph._parent)
    inserted.clear()
    if style:
        inserted.style = style
    if text:
        inserted.add_run(text)
    return inserted


def clear_paragraph(paragraph):
    paragraph.clear()


def find_para(doc, exact):
    for p in doc.paragraphs:
        if " ".join(p.text.split()) == exact:
            return p
    raise ValueError(f"Paragraph not found: {exact}")


def replace_exact(doc, exact, new_text):
    p = find_para(doc, exact)
    clear_paragraph(p)
    p.add_run(new_text)
    return p


def add_block_after(anchor, items):
    current = anchor
    for item in items:
        if isinstance(item, tuple):
            text, style = item
        else:
            text, style = item, None
        current = insert_after(current, text, style=style)
    return current


def main():
    doc = Document(SOURCE)

    replacements = {
        "To design and build an AI-based multi-agent research assistant that automates literature discovery and analysis providing accurate, personalized, and evidence-verified insights to accelerate scientific innovation.":
        "To design and build VeriSci, an AI-based research assistance platform that automates literature discovery, research-gap identification, legal open-access full-text enrichment, knowledge organization, and evidence-grounded question answering using Scopus/Elsevier, IEEE, open academic sources, RAG, and knowledge graph techniques.",

        "Experimental results prove that VeriSci dramatically enhances the efficiency of literature investigation due to its ability to eliminate redundancy, reduce the necessity for human interaction, and improve the relevance of searches. The application of semantic search techniques, the use of graph representation, and the inclusion of conversation artificial intelligence leads to a more sophisticated and engaging research experience. In conclusion, the suggested system shows the possibility of AI-as":
        "Experimental results indicate that VeriSci improves the efficiency of literature investigation by reducing duplicate results, prioritizing verified academic metadata, enabling workspace-based paper saving, and supporting evidence-backed research-gap generation. The use of parallel search, local caching, Scopus verification, legal open-access PDF extraction, RAG-based question answering, and knowledge graph representation creates a more reliable and organized research workflow.",

        "The next stage concentrated on implementing the search and ranking logic. Relevance-scoring algorithms were developed to evaluate the quality of retrieved papers, while evaluation scripts were created to test retrieval accuracy under different query conditions. Once the retrieval system was functioning effectively, the RAG pipeline was introduced. PDF documents downloaded by the system were parsed, cleaned, and transformed into vector embeddings, allowing contextual querying through Large Language Models such as Ollama.":
        "The next stage concentrated on implementing search, ranking, and research-gap discovery. Relevance-scoring algorithms were developed to evaluate retrieved papers using title match, keyword overlap, citation count, publication year, source quality, and Scopus verification. Deep Research functionality was then added to resolve only legal open-access PDFs, download them temporarily, extract useful sections, and delete the temporary files after analysis. The RAG pipeline was integrated for PDF chat using locally stored document chunks and Ollama-based responses.",

        "Step 7: Semantic Chunking and Vector Embedding":
        "Step 7: Semantic Chunking and Local Vector Embedding",

        "The extracted text is divided into smaller, semantically coherent chunks (e.g., by paragraph or section). Each chunk is converted into a high-dimensional vector embedding using a pre-trained language model and indexed into a vector database (e.g., FAISS or ChromaDB) for efficient similarity retrieval.":
        "The extracted text is divided into smaller, semantically coherent chunks. Each chunk is converted into a high-dimensional vector embedding using a sentence-transformer model. Instead of using an external vector database, the current implementation stores document metadata, chunks, and embeddings in a lightweight local pickle-based document store, which keeps the MVP simple and easy to run on a laptop.",

        "Documentation and Research Paper Preparation: The team jointly contributed to report writing, conference paper preparation, plagiarism verification, and project presentations.":
        "Documentation and Research Paper Preparation: The team jointly contributed to report writing, conference paper preparation, result analysis, diagram preparation, and project presentations.",

        "The VeriSci project uses two kinds of databases to handle parts of the system. It uses a Knowledge Graph Database, which is also known as Neo4j to deal with relationships and metadata. The VeriSci project also uses a Custom Local Vector/Document Store for something called Retrieval-Augmented Generation or RAG for short.":
        "The VeriSci project uses three storage mechanisms for different responsibilities. Neo4j is used as a knowledge graph database for relationships between papers, authors, and research fields. A custom local vector/document store is used for Retrieval-Augmented Generation (RAG). In addition, a lightweight SQLite database stores research workspaces and saved paper libraries so that users can continue their literature review work across sessions.",
    }
    for old, new in replacements.items():
        try:
            replace_exact(doc, old, new)
        except ValueError:
            pass

    # Replace the unfinished SOCME placeholder.
    try:
        replace_exact(doc, "bye", "The system-oriented concept map extension (SOCME) for VeriSci represents the project as an interaction between six subsystems: user interaction, literature retrieval, paper access, AI analysis, knowledge organization, and output/storage. The user interaction subsystem begins with the researcher entering a search query, selecting sources, managing a dashboard workspace, and saving relevant papers. The literature retrieval subsystem connects this query to Scopus/Elsevier, IEEE, Semantic Scholar, OpenAlex, and other open academic sources. These sources return metadata such as title, authors, DOI, abstract, publication year, citation count, venue, and source identifiers.")
        anchor = find_para(doc, "The system-oriented concept map extension (SOCME) for VeriSci represents the project as an interaction between six subsystems: user interaction, literature retrieval, paper access, AI analysis, knowledge organization, and output/storage. The user interaction subsystem begins with the researcher entering a search query, selecting sources, managing a dashboard workspace, and saving relevant papers. The literature retrieval subsystem connects this query to Scopus/Elsevier, IEEE, Semantic Scholar, OpenAlex, and other open academic sources. These sources return metadata such as title, authors, DOI, abstract, publication year, citation count, venue, and source identifiers.")
        add_block_after(anchor, [
            "The metadata normalizer converts data from different APIs into a common structure and sends DOI/EID values for verification. The paper access subsystem then checks whether a legal open-access PDF is available through existing PDF links, OpenAlex, Unpaywall, or arXiv. If available, the file is downloaded only to a temporary cache, useful sections are extracted, and the temporary file is deleted after processing. This design avoids heavy storage usage and respects legal access limitations.",
            "The AI analysis subsystem uses extracted metadata, abstracts, snippets, and full-text evidence where available. The RAG pipeline retrieves relevant chunks and passes grounded context to the Ollama-based chatbot or research gap generator. The knowledge organization subsystem creates paper, author, keyword, method, and field relationships in the knowledge graph. Finally, the output and storage subsystem stores selected papers in SQLite workspaces, displays gap reports, supports export, and prepares report-ready outputs.",
            "Figure no. 4.5 SOCME Diagram for VeriSci Research Gap Finder",
        ])
    except ValueError:
        pass

    add_block_after(find_para(doc, "5.3 Important Algorithms"), [
        ("5.3.1 Multi-Source Retrieval Algorithm", None),
        "The multi-source retrieval algorithm accepts a natural language query and forwards it to selected academic sources. Verified sources such as Scopus and IEEE are prioritized for reliable metadata, while open academic sources such as Semantic Scholar, OpenAlex, ArXiv, CrossRef, PubMed, CORE, and Google Scholar extend coverage. The system executes source requests in parallel using worker threads, which reduces waiting time compared with sequential searching.",
        ("Algorithm Steps:", None),
        "1. Accept query, filters, selected sources, year range, citation threshold, and maximum result count.",
        "2. Generate a cache key using the query parameters and return cached results if valid.",
        "3. Dispatch API requests to selected sources in parallel.",
        "4. Normalize all responses into a common paper format.",
        "5. Remove duplicate papers using DOI, ArXiv ID, and fuzzy title similarity.",
        "6. Enrich DOI-bearing papers with Scopus verification where possible.",
        "7. Rank papers using relevance, citation impact, recency, source quality, and metadata completeness.",
        "8. Store the final results temporarily in the session for later analysis, saving, export, or knowledge graph insertion.",
        ("5.3.2 Research Gap Identification Algorithm", None),
        "The research gap algorithm analyzes the selected search results and identifies underexplored areas based on abstracts, keywords, citation patterns, publication years, and available full-text evidence. In normal mode, the system uses metadata and abstracts. In Deep Research mode, it attempts to retrieve legal open-access PDFs and extract compact sections such as abstract, methodology, results, limitations, conclusion, and future work.",
        "The generated gaps are ranked using evidence quality. Papers verified through Scopus, papers with richer abstracts, and papers where legal full text is available receive stronger evidence. This reduces unsupported suggestions and makes the output more useful for final-year project research and literature review work.",
        ("5.3.3 RAG Question Answering Algorithm", None),
        "Uploaded PDFs are parsed using PyPDF2, cleaned, chunked with overlap, embedded using sentence-transformers, and stored in a local document store. When the user asks a question, the query is embedded and compared with existing chunk embeddings using cosine similarity. Maximal Marginal Relevance reranking is used to select diverse and relevant chunks before the context is passed to the LLM. This helps the chatbot answer from document evidence instead of relying only on model memory.",
        ("5.3.4 Knowledge Graph Construction Algorithm", None),
        "The knowledge graph algorithm converts saved or searched paper metadata into graph entities. Paper nodes store title, abstract, year, DOI, citation count, source, journal, and URL details. Author nodes represent contributors, and field nodes represent domains or topics. AUTHORED and IN_FIELD relationships connect these nodes. Indexes on title, author name, and year improve search speed inside Neo4j.",
    ])

    add_block_after(find_para(doc, "5.4 Important Data Structure"), [
        "The project uses multiple data structures because each subsystem has a different storage requirement.",
        "1. Paper Dictionary: The normalized paper dictionary is the core structure used across the system. It contains fields such as title, authors, abstract, year, citations, DOI, EID, source, journal, URL, PDF URL, keywords, subject areas, source metrics, and Scopus verification status.",
        "2. Search Cache JSON: Search responses are cached as JSON files under the data/cache/search directory. The cache key is generated from query parameters, selected sources, and filters. This improves response time for repeated searches.",
        "3. SQLite Workspace Tables: The workspaces table stores local research projects, while the saved_papers table stores selected papers along with notes, tags, source, DOI, year, citations, and the complete paper JSON. A uniqueness constraint prevents duplicate saved papers within the same workspace.",
        "4. RAG Document Store: The local RAG store keeps document metadata, raw text, chunks, and embeddings in a serialized pickle file. This avoids the need for a heavy external vector database during MVP deployment.",
        "5. Neo4j Graph Nodes and Relationships: Paper, Author, and Field nodes are connected using AUTHORED and IN_FIELD relationships. This graph representation supports author exploration, paper search, citation-oriented insights, year distribution, and top field analysis.",
        "6. Deep Research Evidence Object: During Deep Research, each paper is converted into an evidence object containing metadata, evidence level, full-text status, extracted sections, snippets, and text length used for analysis. Temporary PDF files are deleted after evidence extraction.",
    ])

    add_block_after(find_para(doc, "5.5 Graphics User Interface Screenshot"), [
        "The graphical user interface of VeriSci is implemented as a Flask-based web application with a dashboard-oriented layout. The interface provides separate areas for workspace management, academic paper search, knowledge graph exploration, PDF chat, and research gap analysis.",
        "Main UI components include:",
        "1. Dashboard Workspace: Allows users to create and switch research workspaces, view saved paper counts, and continue project-specific literature review work.",
        "2. Search Interface: Allows users to enter a research topic, select verified sources such as Scopus and IEEE, expand to open-access sources, apply year/citation/journal filters, and view ranked paper cards.",
        "3. Paper Detail Panel: Displays abstract, source, DOI, citations, journal, Scopus verification status, PDF availability, and export-related metadata.",
        "4. Research Gap Finder: Provides normal research gap analysis based on metadata and abstracts, and Deep Research gap analysis based on legal open-access full-text extraction where available.",
        "5. Knowledge Graph Explorer: Displays paper relationships and supports searching, statistics, graph visualization, and AI-generated graph insights.",
        "6. PDF Chat Module: Allows PDF upload and document-grounded question answering using the local RAG pipeline.",
        "Suggested screenshots to add in this section: dashboard workspace, search results with Scopus badges, deep research gap output, knowledge graph view, and PDF chat screen.",
    ])

    add_block_after(find_para(doc, "SYSTEM TESTING"), [
        ("7.1 Introduction", None),
        "System testing was performed to verify that the VeriSci platform satisfies its functional and non-functional requirements. Since the system integrates academic APIs, document processing, local storage, RAG, and knowledge graph features, testing was conducted at module level as well as end-to-end workflow level.",
        ("7.2 Testing Objectives", None),
        "The main testing objectives were to verify accurate paper retrieval, correct source filtering, duplicate removal, Scopus verification, legal full-text handling, workspace persistence, PDF upload, RAG response generation, knowledge graph insertion, export functionality, and acceptable response time during repeated searches.",
        ("7.3 Test Cases", None),
        "TC-01 Search papers using Scopus and IEEE: Enter a research query and select verified sources. Expected result: ranked results are returned with title, authors, year, source, DOI/URL, and available abstract.",
        "TC-02 Search using open academic sources: Select Semantic Scholar, OpenAlex, ArXiv, CrossRef, PubMed, CORE, or Google Scholar. Expected result: additional open metadata and PDF links are displayed where available.",
        "TC-03 Duplicate removal: Search a topic that appears in multiple repositories. Expected result: duplicate papers are merged and the richest metadata record is retained.",
        "TC-04 Save paper to workspace: Save a result card into the active workspace. Expected result: paper appears in the saved paper library and persists after page refresh.",
        "TC-05 Deep Research Gaps: Run Deep Research on search results. Expected result: only legal open-access PDFs are downloaded temporarily, text evidence is extracted, and temporary files are deleted after analysis.",
        "TC-06 PDF upload and chat: Upload a readable PDF and ask a question. Expected result: the document is indexed and the chatbot returns an answer grounded in relevant chunks.",
        "TC-07 Knowledge graph insertion: Add search results to the graph. Expected result: Paper, Author, and Field nodes are created and connected correctly.",
        "TC-08 Export results: Export search results as JSON, CSV, or BibTeX. Expected result: downloadable file is generated in the selected format.",
        "TC-09 Missing external service: Run the application when Ollama or Neo4j is not active. Expected result: the application remains usable and displays clear messages for unavailable optional features.",
        ("7.4 Testing Summary", None),
        "Testing showed that the platform can support the main research workflow from query search to paper saving, gap identification, and report preparation. The caching mechanism significantly improves repeated searches. Optional services such as Neo4j and Ollama are handled gracefully, so the application can still perform search and workspace operations even if those services are unavailable.",
    ])

    add_block_after(find_para(doc, "EXPERIMENTAL RESULTS"), [
        ("8.1 Experimental Setup", None),
        "Experiments were conducted on a laptop-based local environment using the Flask web interface. The system was tested with research queries from artificial intelligence, machine learning, natural language processing, and research automation domains. Scopus/Elsevier and IEEE were used as verified sources, while Semantic Scholar, OpenAlex, ArXiv, CrossRef, PubMed, CORE, and Google Scholar were used for wider discovery and open-access coverage.",
        ("8.2 Observed Results", None),
        "The platform successfully retrieved papers from multiple academic repositories, normalized metadata into a common structure, removed duplicate entries, and ranked results according to relevance. The addition of Scopus verification improved trust in search results because the system could distinguish verified indexed records from general open-source metadata. The workspace feature allowed users to save selected papers and continue analysis later without repeating the same search process.",
        "Deep Research mode improved research-gap generation when legal open-access PDFs were available. Instead of relying only on paper titles, the system used abstracts, snippets, and extracted sections such as limitations, methodology, results, conclusion, and future work. This made the gap suggestions more evidence-based. When a legal PDF was not available, the system safely fell back to metadata-only analysis.",
        ("8.3 Performance Observations", None),
        "Parallel source search reduced waiting time compared with sequential API calls. Local caching further reduced latency for repeated queries because previously retrieved results could be returned from the cache instead of calling every external API again. Disabling OCR by default also improved responsiveness, while still allowing OCR to be enabled manually for scanned PDFs.",
        ("8.4 Result Analysis", None),
        "The results indicate that VeriSci is suitable as an MVP for AI-assisted literature review. It combines verified metadata, open-access discovery, temporary full-text evidence extraction, workspace storage, and RAG-based question answering in one interface. The main limitation is that full-text availability depends on legal open-access sources; paywalled papers may provide only metadata and abstracts. Another limitation is that optional services such as Neo4j and Ollama must be installed and running for graph visualization and local chat features.",
        ("8.5 Suggested Quantitative Result Table", None),
        "Add a table with columns: Experiment, Query, Sources Selected, Papers Retrieved, Duplicates Removed, Scopus Verified Papers, Open/PDF Available Papers, First Search Time, Cached Search Time, and Remarks. This will make the result chapter stronger and more measurable.",
    ])

    add_block_after(find_para(doc, "CONCLUSIONS AND FUTURE SCOPE"), [
        ("9.1 Conclusion", None),
        "The VeriSci project successfully demonstrates an AI-based research assistance platform that supports the literature review lifecycle from paper discovery to research-gap identification. The system integrates Scopus/Elsevier, IEEE, and open academic sources to retrieve paper metadata, normalize results, remove duplicates, and rank papers based on relevance and quality. The addition of Scopus verification and source grouping makes the search workflow more professional and reliable.",
        "The platform also improves research productivity through workspace-based storage, legal open-access full-text enrichment, RAG-based PDF chat, and knowledge graph visualization. Deep Research mode addresses a major limitation of title-only search by temporarily downloading legal open-access PDFs, extracting evidence, and deleting temporary files after processing. This keeps storage lightweight while improving the quality of research-gap analysis.",
        "Overall, VeriSci provides a practical MVP for students and researchers who need a unified tool for searching, organizing, analyzing, and understanding research papers. The project shows how retrieval-based AI, local storage, knowledge graphs, and conversational interfaces can be combined to make academic research faster, more structured, and more evidence-driven.",
        ("9.2 Future Scope", None),
        "Future enhancements may include institutional authentication for subscribed full-text sources, improved citation network analysis, automatic literature review generation, stronger evaluation metrics for research-gap quality, live collaboration between multiple users, integration with reference managers such as Zotero or Mendeley, and better graph analytics for identifying emerging research clusters.",
        "The knowledge graph can also be extended with additional node types such as Dataset, Method, Model, Metric, Institution, and Research Gap. This would allow the system to detect not only paper-author-field relationships, but also which methods are commonly used, which datasets are underexplored, and which topics have limited experimental validation.",
        "Further work can include deployment on a cloud platform, authentication and user accounts, report export in DOCX/PDF format, and dashboard analytics showing publication trends, source coverage, citation distribution, and year-wise research growth.",
    ])

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
